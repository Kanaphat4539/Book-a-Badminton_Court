from pathlib import Path
import json
from openpyxl import load_workbook
from openpyxl.styles import Alignment

root=Path.cwd()
book=root/'testing/Test_Case_Book-a-Badminton_Court_Results_20261001_feature-Docx_QA.xlsx'
ev_path=root/'testing/evidence/qa_remaining_probe_results.json'
ev=json.loads(ev_path.read_text(encoding='utf-8'))
updates={
 'TC_CONC_001': ('ผ่าน','Fresh no-cache QA backend image '+ev['environment']['image']+'; isolated MySQL; two unique students sent same Court 1 / 20:00 concurrently -> HTTP 201/400; fixture IDs scoped and removed.'),
 'TC_CONC_019': ('ตรวจไม่ได้','Current-image related malformed-ID probe (POST /bookings/not-an-id/finish) returned HTTP 500; this is not the exact wrong-method/nonexistent-route scenario. Keep TC_CONC_019 unverified; see qa_remaining_probe_results.json.'),
 'TC_ADM_026': ('ผ่าน','TDD fix verified: missing/blank password now raises BadRequestException. Focused/full Jest passed 2/2 and 11/11 suites, 19/19 tests; no-cache backend rebuild image sha256:54ffb880af08d2ec8b563e90f75d90f1ff91e660fe02a30d234e97fc8393e55c; live Admin POST without password returned HTTP400 Password is required; persisted fixture count 0 and Booking count 0. Evidence qa_admin_missing_password_fix.json.'),
 'TC_ADM_028': ('ไม่ผ่าน','Fresh no-cache QA image '+ev['environment']['image']+'; authenticated Admin DELETE student with seeded COMPLETED booking returned HTTP 200; readback proved the booking was deleted (count 1→0) and student was deleted. Contradicts no-silent-history-deletion requirement. Evidence: qa_current_admin_delete_readback.json. Isolated QA fixtures removed.'),
}
wb=load_workbook(book)
for cid,(status,actual) in updates.items():
 matches=[]
 for ws in wb.worksheets[1:10]:
  for row in range(3,ws.max_row+1):
   if ws.cell(row,1).value==cid: matches.append((ws,row))
 if len(matches)!=1: raise ValueError((cid,len(matches)))
 ws,row=matches[0];ws.cell(row,6).value=actual;ws.cell(row,7).value=status;ws.cell(row,6).alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row].height=66
name='21_Fresh_Probe_20261001'
if name in wb.sheetnames: del wb[name]
ws=wb.create_sheet(name)
ws.append(['Case/Probe','Status','Observed actual','Image/source','Fixture cleanup'])
for x in ev['results']:
 ws.append([x['id'],x['status'],x['actual'],ev['environment']['image']+' / '+ev['environment']['commit'],'Unique fixture IDs/bookings removed; final Booking=0, fixture users=0'])
for row in ws.iter_rows():
 for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
for col,width in {'A':32,'B':14,'C':76,'D':66,'E':60}.items(): ws.column_dimensions[col].width=width
wb.save(book)
check=load_workbook(book,read_only=True,data_only=True)
counts={};total=0
for sh in check.worksheets[1:10]:
 for row in range(3,sh.max_row+1):
  cid=sh.cell(row,1).value
  if isinstance(cid,str) and cid.startswith('TC_'):
   total+=1;status=sh.cell(row,7).value;counts[status]=counts.get(status,0)+1
assert total==271,(total,counts)
for cid,(status,_) in updates.items(): assert any(sh.cell(r,1).value==cid and sh.cell(r,7).value==status for sh in check.worksheets[1:10] for r in range(3,sh.max_row+1))
print(json.dumps({'total':total,'counts':counts,'updates':list(updates),'probe_rows':len(ev['results']),'image':ev['environment']['image']},ensure_ascii=False))