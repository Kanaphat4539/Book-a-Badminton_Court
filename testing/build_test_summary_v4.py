"""Rebuild the test summary DOCX from the verified v5 workbook and prior report; preserve source files."""
from pathlib import Path
from collections import Counter
from datetime import datetime
import json
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent
xlsx=root/'Test_Case_Book-a-Badminton_Court_Results_20260928_v5_UI_SEC_CONC.xlsx'
source=root/'Test_Summary_Report_Book-a-Badminton-Court_v3_REQ-BOOK.docx'
out=root/'Test_Summary_Report_Book-a-Badminton-Court_v4_UI_SEC_CONC.docx'
wb=load_workbook(xlsx,read_only=True)
groups={s.title:Counter(s.cell(r,7).value for r in range(3,s.max_row+1)) for s in wb.worksheets[1:10]}
keys=('ผ่าน','ไม่ผ่าน','ตรวจไม่ได้','รอทดสอบ')
total=Counter({k:sum(g[k] for g in groups.values()) for k in keys})
assert total==Counter({'ผ่าน':121,'ไม่ผ่าน':37,'ตรวจไม่ได้':67,'รอทดสอบ':46})
pending=[(s.title,s.cell(r,1).value,str(s.cell(r,2).value or '')) for s in wb.worksheets[1:10] for r in range(3,s.max_row+1) if s.cell(r,7).value=='รอทดสอบ']
assert len(pending)==46 and len(set(p[1] for p in pending))==46
assert all(s.cell(r,6).value in ('',None) for s in wb.worksheets[1:10] for r in range(3,s.max_row+1) if s.cell(r,7).value=='รอทดสอบ')
d=Document(source)

def setp(i,text):
    p=d.paragraphs[i];p.clear();p.add_run(text)

def font_obj(f):
    f.name='TH SarabunPSK'
    rpr=f._element.get_or_add_rPr()
    rf=rpr.rFonts
    if rf is None:
        rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for key in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+key),'TH SarabunPSK')

def format_cell(c,size=11):
    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in c.paragraphs:
        p.paragraph_format.space_after=Pt(0)
        for run in p.runs:run.font.size=Pt(size);font_obj(run.font)

rate=total['ผ่าน']/(total['ผ่าน']+total['ไม่ผ่าน'])*100
setp(2,'เวอร์ชันของระบบที่ทดสอบ: feature-Docx @ 37aea29; ตาราง Test Case: testing/Test_Case_Book-a-Badminton_Court_Results_20260928_v5_UI_SEC_CONC.xlsx')
setp(3,'วันที่จัดทำรายงาน: 28/09/2569 (สรุปผลการทดสอบบนฐาน QA แยก ณ รอบนี้)')
setp(5,'สถานะเอกสาร: รายงานผลทดสอบบางส่วนจาก 271 Test Cases; ไม่ใช่การรับรองความพร้อมใช้งานจริง')
setp(7,f"สถานะภาพรวม: ยังไม่พร้อม Go-Live — ผ่าน {total['ผ่าน']}, ไม่ผ่าน {total['ไม่ผ่าน']}, ตรวจได้บางส่วน/ยังยืนยันไม่ได้ {total['ตรวจไม่ได้']}, รอทดสอบ {total['รอทดสอบ']} จาก 271 เคส. กลุ่ม REQ-UI, REQ-SEC และ REQ-CONC อ้างอิงผลรอบล่าสุดใน workbook v5")
setp(8,'ผล Jest 18 tests / 11 suites และ E2E 1 test เป็นบันทึกจากการทดสอบรอบก่อน ไม่รวมซ้ำใน 271 Test Cases; ผลรอบล่าสุดอ้างอิงหลักฐานราย ID ใน workbook และไฟล์ JSON ที่ระบุไว้')
setp(9,f"อัตราผ่านเฉพาะเคสที่ตัดสินผลได้ = {total['ผ่าน']} ÷ ({total['ผ่าน']} + {total['ไม่ผ่าน']}) = {rate:.1f}% จาก {total['ผ่าน']+total['ไม่ผ่าน']} เคส; ไม่ใช่อัตราผ่านของทุก 271 เคส")
setp(10,'ความเสี่ยงสำคัญ: การสมัครซ้ำกระทบบัญชีเดิม; API Admin ส่ง password hash ใน relation ของ booking (TC_SEC_016); POST จองซ้ำหลายครั้งเมื่อกดยืนยันถี่ (TC_CONC_003); booking id ผิดรูปแบบคืน 500 (TC_CONC_019); UI แยก error/empty state และแสดงข้อมูลการจองบางรายการไม่ครบ')
setp(13,'ขอบเขตที่มีหลักฐาน: รอบก่อนทดสอบ AUTH/BOOK/CANCEL/CHECKIN/BAN/ADMIN ตาม workbook; รอบ 28/09/2569 เพิ่ม browser Chromium, API/MySQL QA และ mocked Date/service สำหรับ REQ-UI (33), REQ-SEC (28), REQ-CONC (23) โดยบางเคสยืนยันได้เพียงบาง assertion')
setp(15,f"ยังไม่ได้ตัดสินผล {total['ตรวจไม่ได้']+total['รอทดสอบ']} เคส: รอทดสอบ {total['รอทดสอบ']} เคส (รายการและเงื่อนไขในหัวข้อ 7.1) และตรวจได้บางส่วน/ยังยืนยันไม่ได้ {total['ตรวจไม่ได้']} เคส (ดู Actual Result ราย ID ใน workbook). ไม่ถือว่าสองกลุ่มนี้ผ่าน")
setp(16,'ข้อจำกัดเวอร์ชัน: ใช้ source commit 37aea29 บน branch feature-Docx; Expected Result ที่มาจาก requirement/revision อื่นต้องตรวจเทียบกับสัญญาพฤติกรรมก่อนตัดสินกรณีข้อความหรือเวลาไม่ตรง')
setp(18,'สภาพแวดล้อม: Windows 11, NestJS/Next.js บน localhost:4001/3001, MySQL QA Docker volume badminton_qa_20260928_mysql_data แยก, Chromium headless; บางกรณีใช้ mocked clock/service. ไม่ใช่ผลของ deployment จริงหรือ Safari/iOS')
setp(19,'หลังทดสอบยืนยันฐาน QA: Booking 0, บัญชี student/admin ส่วนเกิน 0, strikes ของบัญชี seed กลับเป็น 0 และไม่ถูกแบน; services ทดสอบพอร์ต 3001/4001 ถูกปิด. ไม่มีการทดสอบกล้องจริงหรือสภาพแวดล้อม production')
setp(20,'หลักฐานรอบล่าสุด: testing/evidence/qa_ui_suite_*.json, qa_sec_api_focused_results.json, qa_sec_browser_results.json, qa_conc_api_focused_results.json, qa_conc_service_results.json, qa_conc_full_day_results.json และ qa_ui_conc_doubleclick_results.json; ผลราย ID พร้อมข้อจำกัดอยู่ใน workbook v5')
setp(23,'วิธีนับ: ผ่าน = ทุก assertion ในเคสมีหลักฐาน; ไม่ผ่าน = พบข้อขัดกับ Expected Result; ตรวจไม่ได้ = ทดสอบแล้วแต่ยังยืนยันครบไม่ได้; รอทดสอบ = ยังไม่มีหลักฐานการรันตรง ID. อัตราผ่านคิดเฉพาะ ผ่าน + ไม่ผ่าน')
setp(27,'ตาราง Defect เป็นกลุ่มข้อค้นพบเพื่อจัดลำดับแก้ไข ไม่ใช่จำนวน Test Case ที่ไม่ผ่าน: บันทึกเดิม 14 กลุ่ม + ข้อค้นพบรอบ UI/CONC ใหม่ 4 กลุ่ม = 18 กลุ่ม; TC_SEC_016 นับในกลุ่มเดิม ไม่ซ้ำ')
setp(29,'รอบล่าสุด: REQ-UI ผ่าน 3 ไม่ผ่าน 4 ตรวจได้บางส่วน 22 รอ 4; REQ-SEC ผ่าน 14 ไม่ผ่าน 1 ตรวจได้บางส่วน 10 รอ 3; REQ-CONC ผ่าน 10 ไม่ผ่าน 2 ตรวจได้บางส่วน 6 รอ 5. REQ-BOOK รอบก่อนปรับเป็นผ่าน 29 ไม่ผ่าน 10 ตรวจได้บางส่วน 3 (TC_BOOK_037 ทดสอบผ่านแล้ว)')
setp(32,'ความเสี่ยง: API เปิดเผย hash ในข้อมูลการจองแก่ Admin; การกดยืนยันถี่สร้างหลาย POST แม้ฐานข้อมูลคงหนึ่งแถว; id ไม่ใช่เลขคืน 500; Dashboard ขาด error/retry และรายละเอียดหรือสถานะบางรายการ. ไม่สรุปผลเคสที่ยังไม่รันหรือยืนยันไม่ครบว่าใช้งานได้')
setp(33,'งานก่อนเปิดใช้: ตัด sensitive fields ออกจาก response; validate route id ให้คืน 4xx ที่ควบคุมได้; ป้องกัน double-submit; เพิ่ม UI error/retry และตรวจรายละเอียดรายการ; ทดสอบซ้ำเคสไม่ผ่านและ 46 เคสที่รอ พร้อมปิดช่องว่าง 67 เคสที่ตรวจได้บางส่วน')
setp(35,f"ข้อสรุป: ไม่อนุมัติ Go-Live ในรอบนี้ — ยังมี {total['ไม่ผ่าน']} เคสไม่ผ่าน, {total['ตรวจไม่ได้']} เคสยังยืนยันไม่ครบ และ {total['รอทดสอบ']} เคสรอทดสอบ. ลงนามรับรองหลังแก้ไขและมีหลักฐานทดสอบซ้ำครบตามขอบเขต")
setp(37,'Detailed Test Cases (source of truth): testing/Test_Case_Book-a-Badminton_Court_Results_20260928_v5_UI_SEC_CONC.xlsx; 271 ID ใน 9 requirement sheets และ execution sheets สำหรับรอบต่าง ๆ; branch feature-Docx @ 37aea29')
setp(38,'หลักฐาน: testing/evidence/qa_book_*.json, qa_ban_*.json, qa_admin_*.json, qa_ui_suite_*.json, qa_sec_*.json, qa_conc_*.json และ qa_ui_conc_doubleclick_results.json; ชื่อไฟล์เฉพาะและข้อจำกัดแต่ละ ID ดูคอลัมน์ Actual Result ใน workbook')
setp(40,'Performance/Security limitations: ไม่มีรายงาน performance load หลาย instance, ผล security scan, HTTPS deployment จริง, Safari/iOS จริง หรือผล backup/restore ครบ; อย่าตีความผลบน localhost ว่าครอบคลุมสภาพแวดล้อมเหล่านี้')
# Sync every summary table row from workbook values.
for i,(name,stats) in enumerate(groups.items(),1):
    decided=stats['ผ่าน']+stats['ไม่ผ่าน']
    vals=[name,sum(stats.values()),stats['ผ่าน'],stats['ไม่ผ่าน'],stats['ตรวจไม่ได้'],stats['รอทดสอบ'],f"{stats['ผ่าน']/decided*100:.1f}%" if decided else 'N/A']
    for col,val in enumerate(vals):d.tables[0].cell(i,col).text=str(val)
for col,val in enumerate(['รวมทั้งสิ้น (Total)',271,total['ผ่าน'],total['ไม่ผ่าน'],total['ตรวจไม่ได้'],total['รอทดสอบ'],f'{rate:.1f}%']):d.tables[0].cell(10,col).text=str(val)
d.tables[1].cell(3,1).text=f"{total['ผ่าน']} ผ่าน / {total['ไม่ผ่าน']} ไม่ผ่าน / {total['ตรวจไม่ได้']} ตรวจไม่ได้ / {total['รอทดสอบ']} รอทดสอบ"
d.tables[1].cell(3,2).text='อิง workbook v5; ไม่รวม Jest ซ้ำ'
for row in d.tables[1].rows:
    if row.cells[0].text=='REQ-BOOK focused':
        row.cells[1].text='42 เคส: 29 ผ่าน, 10 ไม่ผ่าน, 3 ตรวจบางส่วน, 0 รอ'
        row.cells[2].text='รอบก่อน + TC_BOOK_037 เพิ่ม; ดู QA evidence ใน workbook'
for name,short,ev in [('REQ-UI focused','33 เคส: 3 ผ่าน / 4 ไม่ผ่าน / 22 บางส่วน / 4 รอ','Chromium/UI; qa_ui_suite_*.json'),('REQ-SEC focused','28 เคส: 14 ผ่าน / 1 ไม่ผ่าน / 10 บางส่วน / 3 รอ','API + browser; qa_sec_*.json'),('REQ-CONC focused','23 เคส: 10 ผ่าน / 2 ไม่ผ่าน / 6 บางส่วน / 5 รอ','API/service/UI; qa_conc_*.json และ qa_ui_conc_doubleclick_results.json')]:
    cells=d.tables[1].add_row().cells
    for c,v in zip(cells,(name,short,ev)):c.text=v
sev=[('Critical / Blocker',1,0,1,0),('High',2,0,2,0),('Medium',12,0,12,0),('Low',3,0,3,0),('รวมทั้งสิ้น',18,0,18,0)]
for r,vals in enumerate(sev,1):
    for c,v in enumerate(vals):d.tables[2].cell(r,c).text=str(v)
for entry in [
('QA-UI-10','Medium','TC_UI_018: API /bookings/me 500 แต่ไม่มี error/retry ที่แยกจากรายการว่าง','Open / รอบ v5','เพิ่ม error state และปุ่ม retry'),
('QA-UI-11','Medium','TC_UI_021/023/032: รายละเอียด Court/ID และสถานะ CANCELLED ไม่ครบใน Dashboard ที่ตรวจ','Open / รอบ v5','แสดง court/id/สถานะตามข้อกำหนด; ทดสอบการเรียงรายการ'),
('QA-CONC-12','Medium','TC_CONC_003: DOM.click ยืนยันติดกัน 5 ครั้งทำให้เกิด POST 5 ครั้ง (201,400,400,400,400)','Open / รอบ v5','disable/debounce submit ระหว่าง pending; ทดสอบซ้ำ'),
('QA-CONC-13','Medium','TC_CONC_019: POST /bookings/abc/cancel ตอบ 500 แทน controlled 4xx','Open / รอบ v5','validate numeric id ก่อน query; ทดสอบ route malformed'),
]:
    cells=d.tables[3].add_row().cells
    for c,v in zip(cells,entry):c.text=v
# Explicit 46-ID pending appendix; scenarios come verbatim from XLSX, actions are conservative verification steps.
waiting={
'TC_AUTH_001':'ลงทะเบียนข้อมูลครบ แล้วตรวจ response, DB และ login ต่อเนื่อง',
'TC_AUTH_017':'ลงทะเบียนชื่อไทย/หลายคำ/ช่องว่าง แล้วเทียบข้อมูลที่แสดงและ DB',
'TC_AUTH_019':'ยืนยันกติกาชั้นปี จากนั้นยิงค่า 0/ลบ/ทศนิยม/เกินช่วง',
'TC_AUTH_021':'ยืนยัน password policy แล้วทดสอบความยาว/อักขระ/ข้อความ error',
'TC_AUTH_022':'ยืนยันกติกา case/trim ของ username-email แล้วทดสอบซ้ำกับบัญชี QA',
'TC_AUTH_024':'จำลอง API register ล้มเหลวและ network ขาด; ตรวจ UI ไม่แสดงสำเร็จผิด',
'TC_AUTH_027':'สร้าง 401 หลายคำขอพร้อมกันใน browser; นับ redirect/login',
'TC_AUTH_028':'ลบบัญชี QA หลังออก token; ทดสอบ route ที่ต้องยืนยันตัวตน',
'TC_AUTH_029':'เปิดสองแท็บแล้ว logout หนึ่งแท็บ; ตรวจอีกแท็บไม่ใช้ session เดิม',
'TC_AUTH_031':'บล็อก localStorage แล้วทดสอบ login/redirect/error handling',
'TC_AUTH_032':'สลับบัญชี A/B ใน browser เดียวกัน; ตรวจ cache/ข้อมูลไม่ปน',
'TC_CANC_001':'จอง fixture QA แล้วยกเลิกก่อน deadline; ตรวจ response และ DB',
'TC_CANC_011':'เปิดหน้า cancel และตรวจข้อความเงื่อนไข/เส้นตายที่เห็นจริง',
'TC_CANC_012':'เปิด dialog ยกเลิกและเทียบข้อมูลทุกช่องกับ booking fixture',
'TC_CANC_016':'ตัด network ระหว่าง cancel; ตรวจ toast/status/DB ไม่กล่าวว่าสำเร็จ',
'TC_CANC_017':'เปิดแล้วปิด dialog โดยไม่ยืนยัน; ตรวจไม่มี POST และ DB ไม่เปลี่ยน',
'TC_CHK_001':'ใช้บัญชี Admin QA เปิด Dashboard แล้วตรวจ QR ของแต่ละ court',
'TC_CHK_002':'ใช้กล้อง/QR จริงกับ booking QA ภายใน window; ตรวจ CHECKED_IN',
'TC_CHK_006':'สแกน payload ไม่ใช่เลข court ด้วยกล้องจริง; ตรวจ rejection',
'TC_CHK_007':'สแกน QR ของ booking ที่ CANCELLED; ตรวจไม่เปลี่ยนสถานะ',
'TC_CHK_008':'สแกน QR ซ้ำหลังเช็กอิน; ตรวจ status/response คำขอที่สอง',
'TC_CHK_009':'เปิด Admin ระหว่าง User เช็กอิน; ตรวจการอัปเดตโดยไม่ refresh',
'TC_CHK_010':'หลังเช็กอิน ตรวจ countdown Admin ตามเวลาจริง',
'TC_CHK_011':'ให้ Admin กด finish ก่อนเวลา; ตรวจ status และ DB',
'TC_CHK_013':'ปล่อยครบหนึ่งชั่วโมงให้ scheduler จริงรัน; ตรวจสถานะสุดท้าย',
'TC_CHK_014':'ปฏิเสธ permission กล้องบนอุปกรณ์จริง; ตรวจ fallback/error',
'TC_CHK_020':'ปรับนาฬิกา client เร็ว/ช้า; เทียบ server-side decision',
'TC_CHK_021':'เปิดหน้าสแกนจาก HTTP บนอุปกรณ์จริง; ตรวจ permission/secure context',
'TC_CHK_022':'หน่วง API booking แล้วเปิดหน้าสแกน; ตรวจ state ระหว่างโหลด',
'TC_CHK_023':'ออกจากหน้า/กด Back ระหว่าง scanner active; ตรวจ cleanup กล้อง',
'TC_CHK_024':'จำลอง API check-in fail/network drop; ตรวจ UI/DB ไม่รายงานสำเร็จ',
'TC_CHK_025':'ใช้กล้องหลัง แนวนอน แสงน้อย และ QR เล็กบนอุปกรณ์จริง',
'TC_CHK_026':'ยืนยันกติกาก่อนเริ่มรอบ แล้วสแกน QR ของ Admin ด้วย QA fixture',
'TC_CHK_028':'ควบคุมเวลาไทย 00:00–06:59 ขณะ UTC ยังวันก่อน; ตรวจวัน booking',
'TC_UI_009':'สร้างแจ้งเตือน U10 ทั้งสามกรณี และตรวจข้อความ/เวลาใน UI',
'TC_UI_014':'รวมสอง branch เป้าหมายแล้วตรวจธีมและ layout ของ build นั้น',
'TC_UI_025':'รีเฟรช Dashboard หลังจอง/ยกเลิก/เช็กอินด้วย fixture แต่ละสถานะ',
'TC_UI_033':'จำลอง React render error แล้วตรวจ error boundary/recovery',
'TC_SEC_023':'ตรวจ production deployment config และ secret โดยไม่เปิดเผยค่า',
'TC_SEC_026':'ทดสอบ URL HTTPS ที่ deploy จริงและ mixed-content ใน browser',
'TC_SEC_027':'ยืนยันนโยบาย QR replay/forgery; ทดลอง QR ภาพเก่า/เดาเลขสนาม',
'TC_CONC_013':'รัน backend หลาย instance แยก process แล้วแข่งจอง/cron',
'TC_CONC_014':'snapshot QA แล้ว restart backend+DB; ตรวจข้อมูลไม่สูญ',
'TC_CONC_015':'ทำ backup QA แบบควบคุม; ทดลอง rollback/restore และเทียบแถว',
'TC_CONC_017':'เตรียม MySQL และ SQLite QA แล้วรันทดสอบพฤติกรรมเทียบกัน',
'TC_CONC_018':'รัน happy path ของทุก endpoint ตามรายการ routes และเก็บ response',
}
assert set(waiting)=={cid for _,cid,_ in pending},(set(waiting)-{cid for _,cid,_ in pending},{cid for _,cid,_ in pending}-set(waiting))
signoff=d.paragraphs[41]
p=d.add_paragraph(style='Normal');p.add_run('7.1 หมายเหตุรายการรอทดสอบ (Pending Test Cases)').bold=True
signoff._p.addprevious(p._p)
p=d.add_paragraph('มี 46 เคสสถานะ “รอทดสอบ” ตาม workbook v5; ตารางนี้ไม่ใช่ผลการทดสอบ และไม่ได้เปลี่ยนสถานะเป็นผ่าน. คอลัมน์สุดท้ายคือหลักฐานที่ต้องเก็บก่อนตัดสินผล; หากมีคำว่า “ยืนยันกติกา” ให้ตกลง Expected Result ก่อนรัน')
signoff._p.addprevious(p._p)
t=d.add_table(rows=1,cols=3);t.style='Normal Table';t.alignment=WD_TABLE_ALIGNMENT.CENTER
t.autofit=False
borders=OxmlElement('w:tblBorders')
for edge in ('top','left','bottom','right','insideH','insideV'):
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'C8D3DF');borders.append(el)
t._tbl.tblPr.append(borders)
for i,label in enumerate(('Test Case ID','เรื่องที่รอ','สิ่งที่ต้องตรวจ/เก็บหลักฐาน')):t.rows[0].cells[i].text=label
for group,cid,scenario in pending:
    c=t.add_row().cells
    c[0].text=cid
    c[1].text=scenario
    c[2].text=waiting[cid]
signoff._p.addprevious(t._tbl)
for row in t.rows:
    trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(repeat)
for row in t.rows:
    for cell in row.cells:format_cell(cell,10.5)
# Prevent legacy 8 pt defect rows and keep all document scripts in the requested font.
for sty in d.styles:
    if sty.type in (1,2):font_obj(sty.font)
for p in d.paragraphs:
    for run in p.runs:font_obj(run.font)
for tab in d.tables:
    for row in tab.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for para in c.paragraphs:
                for run in para.runs:
                    font_obj(run.font)
                    if run.font.size is not None and run.font.size.pt < 10:run.font.size=Pt(10)
d.save(out)
r=Document(out)
assert r.styles['Normal'].font.name=='TH SarabunPSK'
assert len(r.tables)==6 and len(r.tables[4].rows)==47
assert r.tables[0].cell(10,2).text=='121' and r.tables[0].cell(10,5).text=='46'
assert all(cid in [row.cells[0].text for row in r.tables[4].rows[1:]] for _,cid,_ in pending)
assert r.tables[3].cell(len(r.tables[3].rows)-1,0).text=='QA-CONC-13'
print(json.dumps({'docx':str(out),'overall':dict(total),'pending_per_group':{k:sum(1 for g,_,__ in pending if g==k) for k in groups},'pending_rows':len(pending),'defect_groups':18,'font':r.styles['Normal'].font.name},ensure_ascii=False))
