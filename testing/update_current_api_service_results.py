from pathlib import Path
import json
from openpyxl import load_workbook
from openpyxl.styles import Alignment
root=Path.cwd(); book=root/'testing/Test_Case_Book-a-Badminton_Court_Results_20261001_feature-Docx_QA.xlsx'; wb=load_workbook(book)
sets=[('qa_ban_service_results.json','Direct Nest service / real isolated MySQL; controlled Date'),('qa_conc_service_results.json','Direct service concurrency / real isolated MySQL; controlled Date'),('qa_admin_api_results.json','HTTP API / current isolated QA backend + MySQL'),('qa_sec_api_focused_results.json','HTTP security API / current isolated QA backend + MySQL')]
updates={}
for filename,method in sets:
 data=json.loads((root/'testing/evidence'/filename).read_text(encoding='utf-8'))
 commit=data.get('environment',{}).get('commit')
 if commit!='60a23c800d56a4f012cdead21adf4ab3f38d0e82': raise AssertionError((filename,commit))
 for item in data.get('results',[]):
  cid=item['id']; status=item['status']; actual=str(item.get('actual',''))
  if status not in ('ผ่าน','ไม่ผ่าน','ตรวจไม่ได้'): raise ValueError((filename,item))
  updates[cid]=(status, f'{actual}\nFresh rerun on HEAD {commit}; QA volume badminton_qa_20260928_mysql_data; evidence testing/evidence/{filename}. Service direct runs used local dist hash verified equal to current image dist.')
for cid,(status,actual) in updates.items():
 found=[]
 for ws in wb.worksheets[1:10]:
  for r in range(3,ws.max_row+1):
   if ws.cell(r,1).value==cid: found.append((ws,r))
 if len(found)!=1: raise AssertionError((cid,len(found)))
 ws,r=found[0];ws.cell(r,6).value=actual;ws.cell(r,7).value=status;ws.cell(r,6).alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[r].height=72
name='22_Current_API_Service_20261001'
if name in wb.sheetnames: del wb[name]
ws=wb.create_sheet(name);ws.append(['Case ID','Status','Fresh observed result','Runner evidence','Source/build'])
for filename,method in sets:
 data=json.loads((root/'testing/evidence'/filename).read_text(encoding='utf-8'))
 for item in data['results']:
  ws.append([item['id'],item['status'],item.get('actual',''),f'testing/evidence/{filename}',f"HEAD {data['environment']['commit']}; isolated QA MySQL; {method}"])
for row in ws.iter_rows():
 for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
for c,w in {'A':23,'B':16,'C':80,'D':52,'E':74}.items():ws.column_dimensions[c].width=w
wb.save(book)
check=load_workbook(book,read_only=True,data_only=True); totals={};count=0
for sh in check.worksheets[1:10]:
 for r in range(3,sh.max_row+1):
  cid=sh.cell(r,1).value
  if isinstance(cid,str) and cid.startswith('TC_'):
   count+=1;st=sh.cell(r,7).value;totals[st]=totals.get(st,0)+1
assert count==271 and sum(totals.values())==271
for cid,(st,_) in updates.items(): assert any(s.cell(r,1).value==cid and s.cell(r,7).value==st for s in check.worksheets[1:10] for r in range(3,s.max_row+1)),cid
assert not [n for n in check.sheetnames if len(n)>31]
print(json.dumps({'mapped_current_results':len(updates),'total':count,'totals':totals,'execution_sheet_rows':check[name].max_row-1},ensure_ascii=False))