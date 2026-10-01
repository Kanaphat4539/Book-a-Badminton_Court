import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
root=Path(__file__).resolve().parents[1]
book=root/'testing/Test_Case_Book-a-Badminton_Court_Results_20261001_feature-Docx_QA.xlsx'
ev=json.loads((root/'testing/evidence/qa_ui_suite_hydration_full_results.json').read_text(encoding='utf-8'))
fix=json.loads((root/'testing/evidence/qa_ui_login_hydration_fix_20261001.json').read_text(encoding='utf-8'))
wb=load_workbook(book)
case_sheets=wb.worksheets[1:10]
updated={}
for item in ev['results']:
    found=[]
    for sh in case_sheets:
        for r in range(3,sh.max_row+1):
            if sh.cell(r,1).value==item['id']: found.append((sh,r))
    if len(found)!=1: raise ValueError((item['id'],len(found)))
    actual=''.join(ch for ch in str(item['actual']) if ch in '\\n\\t' or ord(ch)>=32)[:3000]
    sh,r=found[0]; sh.cell(r,6).value=f"Hydration-fix full UI rerun: {actual}"; sh.cell(r,7).value=item['status']; sh.cell(r,6).alignment=Alignment(wrap_text=True,vertical='top'); updated[item['id']]=item['status']
sm=wb['00_Summary']
sm['A36']='QA feature-Docx: isolated Compose, external QA MySQL; hydration-aware login helper added; no production authentication changes.'
sm['A37']=f"Backend Jest parent-verified 11/11 suites, 18/18 tests; fresh UI rerun {len(ev['results'])} cases: {sum(x['status']=='ผ่าน' for x in ev['results'])} pass / {sum(x['status']=='ไม่ผ่าน' for x in ev['results'])} fail / {sum(x['status']=='ตรวจไม่ได้' for x in ev['results'])} partial/manual. 10/10 real login attempts (5 student, 5 admin) returned HTTP 201 and dashboard."
sm['A38']='Hydration regression reproduced via delayed controlled-input reset; helper retries both fields until values remain stable for 500ms, then asserts successful auth response and dashboard. Do not treat unrelated UI assertion failures as login helper failures.'
sm['A39']='Fresh evidence: testing/evidence/qa_ui_login_hydration_fix_20261001.json; UI suite: testing/evidence/qa_ui_suite_hydration_full_results.json; QA volume badminton_qa_20260928_mysql_data.'
sm['A40']='Source HEAD 60a23c800d56a4f012cdead21adf4ab3f38d0e82 with dirty local changes; no commit/push. Prior 37aea29 evidence remains historical only.'
sm['A41']='Fresh case totals below are recomputed from all requirement rows; UI result statuses reflect current rerun, not old batches.'
sm['A42']='Prior run records remain historical; do not aggregate stale evidence into current status.'
sm['A43']='Older BOOK/BAN/ADMIN results remain attributed to their historical execution unless current evidence updates them.'
sm['A44']='Physical-device/manual-only checks remain partial; DOCX is not finalized.'
name='17_UI_Hydration_20261001'
if name in wb.sheetnames: del wb[name]
s=wb.create_sheet(name); s.append(['Case ID','Area','Method','Status','Observed result'])
for it in ev['results']:
 clean=''.join(ch for ch in str(it['actual']) if ch in '\\n\\t' or ord(ch)>=32)[:3000]
 s.append([it['id'],'User UI','Current isolated QA Chromium; hydration-stable login helper',it['status'],clean])
s.append(['RUN-LOGIN-HELPER','Authentication harness','Delayed controlled-input reset regression test','PASS',fix['regression']['baseline']+'; '+fix['regression']['fixed']])
s.append(['RUN-LOGIN-STUDENT','Authentication UI','Real Chromium form -> application POST -> dashboard','PASS','5/5 attempts HTTP 201 and /dashboard; role STUDENT; no credentials/tokens recorded'])
s.append(['RUN-LOGIN-ADMIN','Authentication UI','Real Chromium form -> application POST -> dashboard','PASS','5/5 attempts HTTP 201 and /dashboard; role ADMIN; no credentials/tokens recorded'])
for c in s[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78')
for row in s.iter_rows(min_row=2):
 for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
 row[0].parent.row_dimensions[row[0].row].height=48
for col,w in {'A':25,'B':26,'C':52,'D':20,'E':100}.items(): s.column_dimensions[col].width=w
s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
wb.save(book)
check=load_workbook(book,read_only=True,data_only=True)
allcases=[]
for sh in check.worksheets[1:10]:
 for r in range(3,sh.max_row+1):
  if str(sh.cell(r,1).value or '').startswith('TC_'): allcases.append((sh.cell(r,1).value,sh.cell(r,7).value))
if len(allcases)!=271: raise AssertionError(f'expected 271 cases, got {len(allcases)}')
counts={}
for _,st in allcases: counts[st]=counts.get(st,0)+1
assert sum(counts.values())==271
for cid,status in updated.items():
 rows=[st for x,st in allcases if x==cid]
 assert rows==[status],(cid,status,rows)
assert check[name].max_row==len(ev['results'])+4
print(json.dumps({'workbook':str(book),'ui_rows_updated':len(updated),'all_requirement_rows':len(allcases),'status_counts':counts,'login_evidence_rows':3},ensure_ascii=False))
