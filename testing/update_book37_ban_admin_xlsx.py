"""Update only XLSX with TC_BOOK_037, REQ-BAN, and REQ-ADMIN fresh evidence."""
from pathlib import Path
from collections import Counter
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

root=Path(r'C:\GuisTee\Project_Y3\Book-a-Badminton\Book-a-Badminton_Court\testing')
ev=root/'evidence'
src=root/'Test_Case_Book-a-Badminton_Court_Results_20260928_v3_REQ-BOOK.xlsx'
out=root/'Test_Case_Book-a-Badminton_Court_Results_20260928_v4_BOOK037_BAN_ADMIN.xlsx'

def load_results(name):
    data=json.loads((ev/name).read_text(encoding='utf-8'))
    return {r['id']:r for r in data['results']}

book37=load_results('qa_book37_ban_ui_results.json')
ban_ui={k:v for k,v in book37.items() if k.startswith('TC_BAN_')}
ban_service=load_results('qa_ban_service_results.json')
ban_real=load_results('qa_ban_real_scheduler_results.json')
ban_restart=load_results('qa_ban_restart_result.json')
admin_ui=load_results('qa_admin_ui_results.json')
admin_ui.update(load_results('qa_admin_ui_recheck_results.json'))
admin_api=load_results('qa_admin_api_results.json')

admin_status={
 1:'ผ่าน',2:'ไม่ผ่าน',3:'ตรวจไม่ได้',4:'ผ่าน',5:'ผ่าน',6:'ตรวจไม่ได้',7:'ผ่าน',8:'ผ่าน',9:'ผ่าน',10:'ไม่ผ่าน',11:'ไม่ผ่าน',12:'ผ่าน',13:'ผ่าน',14:'ผ่าน',15:'ตรวจไม่ได้',16:'ผ่าน',17:'ไม่ผ่าน',18:'ผ่าน',19:'ไม่ผ่าน',20:'ผ่าน',21:'ไม่ผ่าน',22:'ไม่ผ่าน',23:'ผ่าน',24:'ผ่าน',25:'ผ่าน',26:'ไม่ผ่าน',27:'ผ่าน',28:'ไม่ผ่าน',29:'ผ่าน',30:'ผ่าน',31:'ไม่ผ่าน',32:'ผ่าน',33:'ผ่าน',34:'ผ่าน',35:'ผ่าน',36:'ผ่าน',37:'ผ่าน',38:'ผ่าน',39:'ตรวจไม่ได้',40:'ตรวจไม่ได้',41:'ไม่ผ่าน',42:'ตรวจไม่ได้'
}
admin_status={f'TC_ADM_{i:03d}':s for i,s in admin_status.items()}
assert Counter(admin_status.values())==Counter({'ผ่าน':25,'ไม่ผ่าน':11,'ตรวจไม่ได้':6})
ban_status={f'TC_BAN_{i:03d}':'ผ่าน' for i in range(1,27)}

w=load_workbook(src)

def rows_for(sheet):
    return {sheet.cell(r,1).value:r for r in range(3,sheet.max_row+1)}

def apply(sheet_name,status_map,actual_map,evidence_note):
    sh=w[sheet_name];rows=rows_for(sh)
    assert set(status_map)<=set(rows)
    for cid,status in status_map.items():
        r=rows[cid];sh.cell(r,7).value=status
        if cid in actual_map:
            sh.cell(r,6).value=f"{actual_map[cid]['actual']}\nหลักฐาน: {evidence_note}"
    return rows

book_rows=apply('02_REQ-BOOK',{'TC_BOOK_037':'ผ่าน'},{'TC_BOOK_037':book37['TC_BOOK_037']},'testing/evidence/qa_book37_ban_ui_results.json')

ban_actual={}
for source in (ban_ui,ban_service,ban_real,ban_restart):ban_actual.update(source)
ban_rows=apply('05_REQ-BAN',ban_status,ban_actual,'testing/evidence/qa_ban_real_scheduler_results.json; qa_ban_service_results.json; qa_ban_restart_result.json; qa_book37_ban_ui_results.json; หลักฐานเดิมในเซลล์สำหรับ TC_BAN_003/006-011/013/017-019')

admin_actual=dict(admin_ui)
# API evidence is authoritative for cases whose full assertions are API/data/concurrency focused.
for cid in ['TC_ADM_016','TC_ADM_019','TC_ADM_020','TC_ADM_021','TC_ADM_024','TC_ADM_026','TC_ADM_028','TC_ADM_029','TC_ADM_033']:
    admin_actual[cid]=admin_api[cid]
admin_rows=apply('06_REQ-ADMIN',admin_status,admin_actual,'testing/evidence/qa_admin_ui_results.json; qa_admin_ui_recheck_results.json; qa_admin_api_results.json')

# Keep the existing BOOK execution trace in sync with the focused recheck.
ex=w['14_REQ-BOOK_Execution']
for r in range(2,ex.max_row+1):
    if ex.cell(r,1).value=='TC_BOOK_037':
        ex.cell(r,2).value='ผ่าน';ex.cell(r,3).value=w['02_REQ-BOOK'].cell(book_rows['TC_BOOK_037'],6).value;ex.cell(r,4).value='testing/evidence/qa_book37_ban_ui_results.json';ex.cell(r,5).value='feature-Docx 37aea29';break
else:raise AssertionError('TC_BOOK_037 missing in execution sheet')

def execution_sheet(title,source_sheet,status_map,evidence_set):
    if title in w.sheetnames:del w[title]
    sh=w.create_sheet(title);sh.append(['Test Case ID','Status','Actual Result','Evidence Set','Build'])
    srcsh=w[source_sheet];srcrows=rows_for(srcsh)
    for cid in sorted(status_map):sh.append([cid,status_map[cid],srcsh.cell(srcrows[cid],6).value,evidence_set,'feature-Docx 37aea29'])
    for c in sh[1]:c.fill=PatternFill('solid',fgColor='DCE6F1');c.font=Font(name='Tahoma',size=10,bold=True)
    for row in sh.iter_rows(min_row=2):
        for c in row:c.font=Font(name='Tahoma',size=9);c.alignment=Alignment(wrap_text=True,vertical='top')
        sh.row_dimensions[row[0].row].height=82
    for col,width in {'A':20,'B':16,'C':105,'D':52,'E':28}.items():sh.column_dimensions[col].width=width
    sh.freeze_panes='A2';sh.auto_filter.ref=f'A1:E{sh.max_row}'

execution_sheet('15_REQ-BAN_Execution','05_REQ-BAN',ban_status,'qa_ban_*.json; qa_book37_ban_ui_results.json')
execution_sheet('16_REQ-ADMIN_Execution','06_REQ-ADMIN',admin_status,'qa_admin_api_results.json; qa_admin_ui*.json')

case_sheets=w.worksheets[1:10]
by_group={sh.title:Counter(sh.cell(r,7).value for r in range(3,sh.max_row+1)) for sh in case_sheets}
keys=['ผ่าน','ไม่ผ่าน','ตรวจไม่ได้','รอทดสอบ']
counts=Counter({k:sum(group[k] for group in by_group.values()) for k in keys})
assert by_group['02_REQ-BOOK']==Counter({'ผ่าน':29,'ไม่ผ่าน':10,'ตรวจไม่ได้':3})
assert by_group['05_REQ-BAN']==Counter({'ผ่าน':26})
assert by_group['06_REQ-ADMIN']==Counter({'ผ่าน':25,'ไม่ผ่าน':11,'ตรวจไม่ได้':6})
assert counts==Counter({'ผ่าน':111,'ไม่ผ่าน':31,'ตรวจไม่ได้':38,'รอทดสอบ':91})
assert sum(counts.values())==271

summary=w['00_Summary']
summary.cell(37,1).value='ผลทดสอบจริง 28/09/2569: เพิ่ม TC_BOOK_037 และรอบโฟกัส REQ-BAN/REQ-ADMIN บน Chromium, API, MySQL QA แยก, scheduler จริงหลัง restart และ service/cron ด้วย mocked Date; อัปเดตเฉพาะ XLSX ตามคำสั่งผู้ใช้'
for col,key in [(2,'ผ่าน'),(4,'ไม่ผ่าน'),(6,'ตรวจไม่ได้'),(8,'รอทดสอบ')]:summary.cell(38,col).value=counts[key]
summary.cell(39,1).value='ข้อจำกัดรอบนี้: REQ-ADMIN ยังตรวจไม่ได้ 6 เคส (CI/ดีไซน์, flow กล้องซ้ำ, read/update policy, guest news policy); เคสกลุ่ม Requirement อื่นที่ยังรอหรือ partial ไม่ได้เปลี่ยนสถานะ'
summary.cell(40,1).value='หลักฐานรอบนี้: testing/evidence/qa_book37_ban_ui_results.json, qa_ban_*.json, qa_admin_api_results.json, qa_admin_ui_results.json และ qa_admin_ui_recheck_results.json'
summary.cell(42,1).value='รอบ REQ-BOOK ล่าสุด: TC_BOOK_037 ผ่านแล้ว; REQ-BOOK รวม 42 เคส — ผ่าน 29, ไม่ผ่าน 10, ตรวจบางส่วน 3, รอทดสอบ 0'
summary.append(['รอบ BOOK037/BAN/ADMIN 28/09/2569: BOOK037 ผ่าน; REQ-BAN 26/26 ผ่าน; REQ-ADMIN ผ่าน 25, ไม่ผ่าน 11, ตรวจไม่ได้ 6; ผลรวมทั้ง workbook ผ่าน 111, ไม่ผ่าน 31, ตรวจไม่ได้ 38, รอทดสอบ 91 จาก 271'])
summary.row_dimensions[summary.max_row].height=58
summary.cell(summary.max_row,1).alignment=Alignment(wrap_text=True,vertical='top')
summary.cell(summary.max_row,1).font=Font(name='Tahoma',size=10,bold=True)
w.calculation.fullCalcOnLoad=True
w.save(out)

# Reopen and verify values rather than trusting save.
check=load_workbook(out,read_only=True)
check_groups={sh.title:Counter(sh.cell(r,7).value for r in range(3,sh.max_row+1)) for sh in check.worksheets[1:10]}
check_counts=Counter({k:sum(g[k] for g in check_groups.values()) for k in keys})
assert check_counts==counts
assert check['00_Summary'].cell(38,2).value==111
assert check['00_Summary'].cell(38,4).value==31
assert check['00_Summary'].cell(38,6).value==38
assert check['00_Summary'].cell(38,8).value==91
assert check['15_REQ-BAN_Execution'].max_row==27
assert check['16_REQ-ADMIN_Execution'].max_row==43
print(json.dumps({'xlsx':str(out),'overall':dict(counts),'groups':{k:dict(v) for k,v in by_group.items() if k in ('02_REQ-BOOK','05_REQ-BAN','06_REQ-ADMIN')},'docx_updated':False},ensure_ascii=False))
