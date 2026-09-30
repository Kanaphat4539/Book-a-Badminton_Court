"""Merge only fresh, attributable UI/SEC/CONC QA observations into a NEW XLSX; never touch DOCX."""
from pathlib import Path
from collections import Counter
from math import ceil
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
root=Path(r'C:\GuisTee\Project_Y3\Book-a-Badminton\Book-a-Badminton_Court\testing')
ev=root/'evidence'
src=root/'Test_Case_Book-a-Badminton_Court_Results_20260928_v4_BOOK037_BAN_ADMIN.xlsx'
out=root/'Test_Case_Book-a-Badminton_Court_Results_20260928_v5_UI_SEC_CONC.xlsx'

def read(name):
    data=json.loads((ev/name).read_text(encoding='utf-8'))
    assert data['environment']['commit']=='37aea29',name
    assert data['environment']['dbVolume']=='badminton_qa_20260928_mysql_data',name
    return {r['id']:dict(r,evidence=name) for r in data['results']}
ui={}
for f in ['qa_ui_suite_a_results.json','qa_ui_suite_b_results.json','qa_ui_suite_ui_recheck_results.json','qa_ui_suite_ui_focus_results.json','qa_ui_suite_ui_last_results.json','qa_ui_suite_ui_exact_results.json']:
    ui.update(read(f))
sec={}
for f in ['qa_sec_api_focused_results.json','qa_sec_browser_results.json']:
    sec.update(read(f))
conc={}
for f in ['qa_conc_api_focused_results.json','qa_conc_service_results.json','qa_conc_full_day_results.json','qa_ui_conc_doubleclick_results.json']:
    conc.update(read(f))
assert len(ui)>=30 and len(sec)>=25 and len(conc)>=17,(len(ui),len(sec),len(conc))
# Tests run with a narrow assertion are partial, even if the runner's narrow assertion passed.
# A runner timeout/selector error is NOT an application defect; leave those cases pending.
ui_states={
 1:'ตรวจไม่ได้',2:'ตรวจไม่ได้',3:'ตรวจไม่ได้',4:'ผ่าน',5:'ผ่าน',6:'ตรวจไม่ได้',7:'ตรวจไม่ได้',8:'ตรวจไม่ได้',10:'ตรวจไม่ได้',11:'ตรวจไม่ได้',12:'ตรวจไม่ได้',13:'ผ่าน',15:'ตรวจไม่ได้',16:'ตรวจไม่ได้',17:'ตรวจไม่ได้',18:'ไม่ผ่าน',19:'ตรวจไม่ได้',20:'ตรวจไม่ได้',21:'ไม่ผ่าน',22:'ตรวจไม่ได้',23:'ไม่ผ่าน',24:'ตรวจไม่ได้',26:'ตรวจไม่ได้',27:'ตรวจไม่ได้',28:'ตรวจไม่ได้',29:'ตรวจไม่ได้',30:'ตรวจไม่ได้',31:'ตรวจไม่ได้',32:'ไม่ผ่าน'
}
sec_states={1:'ผ่าน',2:'ผ่าน',3:'ผ่าน',4:'ผ่าน',5:'ผ่าน',6:'ผ่าน',7:'ผ่าน',8:'ผ่าน',9:'ผ่าน',10:'ตรวจไม่ได้',11:'ผ่าน',12:'ตรวจไม่ได้',13:'ตรวจไม่ได้',14:'ตรวจไม่ได้',15:'ตรวจไม่ได้',16:'ไม่ผ่าน',17:'ผ่าน',18:'ผ่าน',19:'ตรวจไม่ได้',20:'ผ่าน',21:'ตรวจไม่ได้',22:'ผ่าน',24:'ตรวจไม่ได้',25:'ตรวจไม่ได้',28:'ตรวจไม่ได้'}
conc_states={1:'ผ่าน',2:'ผ่าน',3:'ไม่ผ่าน',4:'ผ่าน',5:'ผ่าน',6:'ตรวจไม่ได้',7:'ผ่าน',8:'ตรวจไม่ได้',9:'ผ่าน',10:'ผ่าน',11:'ตรวจไม่ได้',12:'ตรวจไม่ได้',16:'ผ่าน',19:'ไม่ผ่าน',20:'ตรวจไม่ได้',21:'ผ่าน',22:'ผ่าน',23:'ตรวจไม่ได้'}
status_sets={
 '07_REQ-UI':{f'TC_UI_{n:03d}':v for n,v in ui_states.items()},
 '08_REQ-SEC':{f'TC_SEC_{n:03d}':v for n,v in sec_states.items()},
 '09_REQ-CONC':{f'TC_CONC_{n:03d}':v for n,v in conc_states.items()},
}
evidence={'07_REQ-UI':ui,'08_REQ-SEC':sec,'09_REQ-CONC':conc}
notes={
 'TC_UI_003':'ตรวจเฉพาะข้อความวันนี้/ไม่มี input วันที่; ยังไม่ได้เทียบเลขวันที่และ timezone อื่น',
 'TC_UI_005':'mocked browser clock 28/09/2569 เวลา 14:00 เฉพาะการแสดงผล; ไม่มี POST จอง',
 'TC_UI_006':'ทดสอบสลับสี User pages; Admin และ login/register ไม่ครบ',
 'TC_UI_007':'ธีม Dark คงอยู่หลัง reload และไป booking; flash ช่วงโหลดไม่ได้วัด',
 'TC_UI_010':'เห็นเมนูล่าง 4 รายการบนจอมือถือ; ไม่ได้กดเข้าทีละเมนู',
 'TC_UI_015':'ตรวจ page-level overflow 5 widths; ไม่ได้พิสูจน์ปุ่ม/ข้อความถูกบังทั้งหมด',
 'TC_UI_017':'หน้า dashboard fixture ว่างไม่พบชื่อเดือน กันยายน; ยังไม่ได้เทียบทุกหน้า/ทุก timezone จึงยังไม่ตัดสิน Fail',
 'TC_UI_018':'กรณี API /bookings/me จำลอง HTTP 500 ไม่พบข้อความ error/ปุ่ม retry; empty state กับ error แยกไม่ชัด',
 'TC_UI_020':'บล็อก external hosts แล้ว login/booking ยังใช้งานได้; ไม่ได้ตรวจ font glyph/alt ทุกจุด',
 'TC_UI_021':'Fixture Court 3/18:00/PENDING: Recent Bookings แสดงเพียง Court, PENDING, 18:00, วันที่ ISO; ไม่เห็น Court 3 หรือ booking ID ตาม Expected',
 'TC_UI_023':'มี PENDING,CHECKED_IN,COMPLETED,CANCELLED ใน QA แต่หน้า Dashboard แสดงข้อความสามสถานะแรก ไม่พบ CANCELLED; Recent Bookings ใช้ slice(0,3)',
 'TC_UI_024':'ตรวจเพียงไม่มี NaN/เวลาติดลบ; ไม่ได้คุมเวลาข้าม 11:00',
 'TC_UI_026':'ชื่อไทยยาว/จอ 320px ไม่ล้น; ยังไม่ได้ทดสอบ major/admin ว่างตามขั้นตอน',
 'TC_UI_027':'ตรวจ route เฉพาะ Student; Admin route matrix ยังไม่ได้ตรวจ',
 'TC_UI_030':'Dialog จอง: focus ภายใน, Escape ปิด, focus กลับ opener; Dialog ยกเลิกยังไม่ทดสอบ',
 'TC_UI_032':'Fixture 4 สถานะ แต่ Dashboard เห็นข้อความ 3 สถานะ ไม่พบ CANCELLED; screen reader/loading status ยังไม่ได้ตรวจ',
 'TC_SEC_010':'API login ไม่มี password/hash; register response/console/backend logs ยังไม่ตรวจ',
 'TC_SEC_012':'mock protected 401 ล้าง token/user; observed 2 navigation events to /login จึงยังยืนยัน redirect เพียงครั้งเดียวไม่ได้',
 'TC_SEC_015':'fetch ตรง ไม่ผ่าน Axios interceptor จึงยังพิสูจน์คำขอเก่าของ client จริงไม่ได้',
 'TC_SEC_016':'GET /bookings และ /bookings/notifications มี nested password key รั่วจริง; /bookings/me ไม่พบ',
 'TC_SEC_019':'แก้ localStorage role/id แล้วยิง /users ยัง 403; การส่ง stu_id ของ B ใน POST booking ไม่สำเร็จเนื่องจากเวลาทดสอบ จึงยังไม่ครบ',
 'TC_SEC_024':'ลอง login ผิด 20 ครั้งใน QA ได้ 401 ทั้งหมด ผู้ใช้อื่นยังใช้งาน; นโยบายจำกัดความพยายามยังไม่กำหนด',
 'TC_SEC_028':'logout ล้าง localStorage แต่ JWT เดิม replay GET /bookings/me ยังได้ 200; ต้องยืนยัน revocation policy ก่อนตัดสิน',
 'TC_CONC_003':'คลิกยืนยัน 5 ครั้งติดกันด้วย DOM.click: frontend ยิง POST 5 ครั้ง (201,400,400,400,400) แม้ DB สร้าง 1 รายการ; ไม่ผ่านข้อกำหนดตัวกันการส่งซ้ำฝั่งหน้าจอ',
 'TC_CONC_006':'cron พร้อมกัน 3 รอบใน process เดียวให้ strike=3 ตาม 3 bookings; ยังไม่ทดสอบข้าม process',
 'TC_CONC_008':'60 court-hour slots สถานะ COMPLETED ครบ, availability 60 รายการ, notification API ตรงข้อความ; ยังไม่ได้ตรวจการแสดงบน Dashboard',
 'TC_CONC_011':'finish+cron พร้อมกันใน process เดียวหนึ่งรอบ final COMPLETED; ยังไม่ยืนยันทุก interleaving/ประวัติและโควตา',
 'TC_CONC_012':'cron เพิ่ม strike ที่สอง, แบน 24h และคำขอหลังแบน 400; ยังไม่ได้แข่ง createBooking พร้อม cron จริง',
 'TC_CONC_019':'PUT /auth/login=404, DELETE /courts=404, POST /bookings/abc/cancel=500, unknown route=404; id ไม่ใช่เลขทำให้ TypeORM query NaN/SQL error ไม่ใช่ controlled 4xx',
 'TC_CONC_020':'ลอง field เกินตอนจองกับผู้ใช้ quota เต็ม ได้ 400 ก่อนพิสูจน์ field ถูก ignore; สมัคร field เกินเดิมไม่ได้ rerun',
 'TC_CONC_023':'local OPTIONS Origin untrusted ได้ Access-Control-Allow-Origin=*; ยังไม่ทราบ deployed origin/policy จึงไม่ตัดสินผ่าน/ไม่ผ่าน',
}
wb=load_workbook(src)
original={}
for sheet_name,stat in status_sets.items():
    sh=wb[sheet_name]
    ids={sh.cell(r,1).value:r for r in range(3,sh.max_row+1)}
    assert len(ids)==sh.max_row-2 and set(stat)<=set(ids)
    original[sheet_name]=Counter(sh.cell(r,7).value for r in ids.values())
    for cid,new in stat.items():
        assert cid in evidence[sheet_name],f'{cid} lacks fresh run evidence'
        e=evidence[sheet_name][cid];r=ids[cid]
        # Never turn a tool/selector failure into an app failure.
        if new=='ไม่ผ่าน':
            assert cid in {'TC_UI_018','TC_UI_021','TC_UI_023','TC_UI_032','TC_SEC_016','TC_CONC_003','TC_CONC_019'},cid
        sh.cell(r,7).value=new
        detail=notes.get(cid,e['actual'])
        sh.cell(r,6).value=f"{detail}\nหลักฐาน: testing/evidence/{e['evidence']} (QA MySQL แยก; source commit 37aea29)"
    title={'07_REQ-UI':'17_REQ-UI_Execution','08_REQ-SEC':'18_REQ-SEC_Execution','09_REQ-CONC':'19_REQ-CONC_Execution'}[sheet_name]
    if title in wb:del wb[title]
    x=wb.create_sheet(title)
    x.append(['Test Case ID','Status','Actual Result','Evidence File','Build'])
    for cid in ids:
        r=ids[cid];er=evidence[sheet_name].get(cid)
        f='testing/evidence/'+er['evidence'] if cid in stat and er else '— (รอทดสอบรอบนี้)'
        x.append([cid,sh.cell(r,7).value,sh.cell(r,6).value,f,'feature-Docx 37aea29'])
    for c in x[1]:c.fill=PatternFill('solid',fgColor='DCE6F1');c.font=Font(name='Tahoma',size=10,bold=True)
    for row in x.iter_rows(min_row=2):
        for c in row:c.font=Font(name='Tahoma',size=9);c.alignment=Alignment(wrap_text=True,vertical='top')
        x.row_dimensions[row[0].row].height=min(240,max(48,ceil(len(str(row[2].value or ''))/78)*18))
    for col,width in {'A':20,'B':16,'C':100,'D':48,'E':26}.items():x.column_dimensions[col].width=width
    x.freeze_panes='A2';x.auto_filter.ref=f'A1:E{x.max_row}'
keys=['ผ่าน','ไม่ผ่าน','ตรวจไม่ได้','รอทดสอบ']
groups={sh.title:Counter(sh.cell(r,7).value for r in range(3,sh.max_row+1)) for sh in wb.worksheets[1:10]}
overall=Counter({k:sum(g[k] for g in groups.values()) for k in keys})
assert sum(overall.values())==271
s=wb['00_Summary']
s.cell(37,1).value='ผลทดสอบจริง 28/09/2569: เพิ่ม REQ-UI/SEC/CONC บน Chromium/API/MySQL QA แยกและ service mocked Date; หลักฐานเฉพาะเคสที่รันครบ/บางส่วนเท่านั้น; DOCX ยังไม่ได้อัปเดต'
for col,key in [(2,'ผ่าน'),(4,'ไม่ผ่าน'),(6,'ตรวจไม่ได้'),(8,'รอทดสอบ')]:s.cell(38,col).value=overall[key]
s.cell(39,1).value='ข้อจำกัด: มือถือ Safari จริง, deployment HTTPS/CORS, production config, backup/restore, multi-instance, และ UI flow บางกรณีไม่ได้ตรวจครบ; เคสเหล่านั้นคงรอทดสอบหรือตรวจไม่ได้ตามหลักฐาน'
s.cell(40,1).value='หลักฐานเพิ่ม: testing/evidence/qa_ui_suite_*.json, qa_sec_api_focused_results.json, qa_sec_browser_results.json, qa_conc_api_focused_results.json, qa_conc_service_results.json, qa_conc_full_day_results.json, qa_ui_conc_doubleclick_results.json'
s.cell(42,1).value='ผลรอบก่อน (ก่อนทดสอบ REQ-UI/SEC/CONC ครั้งนี้): REQ-BOOK 42 เคส — ผ่าน 29 ไม่ผ่าน 10 ตรวจบางส่วน 3; ตัวเลขรวมในบรรทัดเดิมเป็นยอดประวัติ ให้ยึดยอดล่าสุดที่แถว 38 และท้ายชีต'
s.append([f"รอบ UI/SEC/CONC 28/09/2569: รวม 271 ผ่าน {overall['ผ่าน']} ไม่ผ่าน {overall['ไม่ผ่าน']} ตรวจไม่ได้ {overall['ตรวจไม่ได้']} รอ {overall['รอทดสอบ']}; เฉพาะ XLSX รุ่น v5, DOCX รอผู้ใช้สั่งหยุด"])
s.row_dimensions[s.max_row].height=52;s.cell(s.max_row,1).alignment=Alignment(wrap_text=True,vertical='top')
wb.calculation.fullCalcOnLoad=True
wb.save(out)
check=load_workbook(out,read_only=True)
assert len(check.worksheets)==20
assert check['00_Summary'].cell(38,2).value==overall['ผ่าน']
assert check['00_Summary'].cell(38,4).value==overall['ไม่ผ่าน']
assert check['00_Summary'].cell(38,6).value==overall['ตรวจไม่ได้']
assert check['00_Summary'].cell(38,8).value==overall['รอทดสอบ']
assert sum(Counter(sh.cell(r,7).value for sh in check.worksheets[1:10] for r in range(3,sh.max_row+1)).values())==271
print(json.dumps({'xlsx':str(out),'overall':dict(overall),'groups':{k:dict(groups[k]) for k in status_sets},'prior':{k:dict(v) for k,v in original.items()},'docx_updated':False},ensure_ascii=False))
