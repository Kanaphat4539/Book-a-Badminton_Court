import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment
root=Path('.')
evidence_paths=sorted((root/'testing/evidence').glob('qa_pending_api_batch_*.json'), key=lambda p:p.stat().st_mtime)
evidence=json.loads(evidence_paths[-1].read_text(encoding='utf-8'))
evidence_file=str(evidence_paths[-1]).replace('\\\\','/')
path=root/'testing/Test_Case_Book-a-Badminton_Court_Results_20261001_feature-Docx_QA.xlsx'
wb=load_workbook(path)
by_id={r['id']:r for r in evidence['results']}
updated=[]
for target_sheet in ['01_REQ-AUTH','02_REQ-BOOK','06_REQ-ADMIN']:
    ws=wb[target_sheet]
    for row in ws.iter_rows(min_row=1):
        case_id=row[0].value if row else None
        if case_id in by_id:
            r=by_id[case_id]
            ws.cell(row=row[0].row,column=6).value=f"Live API batch {evidence['environment']['started']}; HTTP/API result: {r['actual']}\\nEvidence: {evidence_file}"
            ws.cell(row=row[0].row,column=7).value='ผ่าน' if r['status']=='PASS' else 'ไม่ผ่าน'
            updated.append(case_id)
assert set(updated)==set(by_id), (updated, list(by_id))
name='25_Pending_API_20261001'
if name in wb.sheetnames: del wb[name]
ws=wb.create_sheet(name)
ws.append(['Case ID','Result','Observed result','Evidence file','Source commit','Backend image','Database volume','Cleanup'])
for r in evidence['results']:
    ws.append([r['id'],r['status'],r['actual'],evidence_file,evidence['environment']['head'],evidence['environment']['backendImage'],evidence['environment']['databaseVolume'],json.dumps(evidence['environment']['cleanup'])])
for col,width in {'A':22,'B':16,'C':95,'D':68,'E':45,'F':70,'G':40,'H':60}.items(): ws.column_dimensions[col].width=width
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
wb.save(path)
print(json.dumps({'workbook':str(path),'updated_ids':updated,'count':len(updated),'summary':evidence['summary'],'sheet':name},ensure_ascii=False))