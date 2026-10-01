from pathlib import Path
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Alignment

path = Path('testing/Test_Case_Book-a-Badminton_Court_Results_20261001_feature-Docx_QA.xlsx')
wb = load_workbook(path)
name = '15_Fresh_Run_Evidence'
if name in wb.sheetnames:
    del wb[name]
ws = wb.create_sheet(name)
rows = [
    ['Check ID', 'Scope', 'Command / Method', 'Result', 'Fresh observation'],
    ['RUN-AUTH-DIRECT', 'Current checkout MySQL API login reproducer', 'Docker QA current build; POST /auth/login with isolated seeded test user; inspect response keys only', 'PASS — DIRECT API', 'HTTP 201; response keys access_token,user. This does not prove browser redirect or every user role. No SQL error reproduced.'],
    ['RUN-AUTH-BROWSER', 'Historical browser login timeout', 'Earlier Playwright waitForURL /dashboard; source fingerprint not reliably attributable', 'UNVERIFIED — NOT A CURRENT CASE RESULT', 'Old backend log excerpt showed 1064 near `)) LIMIT 1` but full generated SQL/binds are absent. Current direct API succeeds. Cause unknown; do not label as app defect.'],
    ['RUN-CONC-REG', 'Registration concurrency', 'Delegated isolated .qa-deps SQL.js run', 'PASS — HISTORICAL FRESH RUN', 'Earlier task recorded one concurrent registration accepted and duplicate returned 400. This run record is transcript-backed; not rerun in this continuation.'],
    ['RUN-MYSQL-API', 'MySQL/API cases', 'Previous delegated runners, isolated Docker QA volume', 'MIXED — AGGREGATE ONLY', '49 aggregate: 26 pass, 20 fail, 3 blocked/partial; no attributable ID-to-result map for these counts in this workbook. Do not assign them to individual case statuses or characterize all failures as app bugs.'],
    ['RUN-UI-001', 'Browser responsive layout', 'Previous Playwright batch, dashboard/booking 1440x900 and 1920x1080', 'PASS — HISTORICAL FRESH RUN', 'Transcript records no horizontal overflow and no header/footer/main overlap. Current continuation did not rerun browser suite.'],
    ['RUN-UI-003', 'Browser date constraint', 'Previous Playwright batch', 'PASS — HISTORICAL FRESH RUN', 'Transcript records today-only hint and no date input. Current continuation did not rerun browser suite.'],
    ['RUN-UI-OTHER', 'Remaining browser suite', 'No attributable fresh full run', 'BLOCKED / NOT COUNTED', 'Login-dependent flows were not considered passed.'],
    ['RUN-QA-CLEANUP', 'QA fixture/container cleanup', 'docker compose down; verify named volume remains', 'PASS', 'QA containers stopped; badminton_qa_20260928_mysql_data remains. During direct auth reproduction the isolated Booking table count was zero; no production service/data was used.'],
    ['RUN-UNIT-AUTH', 'AuthService lookup boundary regression', 'cd backend && npm test -- --runInBand src/auth/auth.service.spec.ts', 'PASS', '1 suite / 1 test passed; verifies username is sent through findByUsername and absent user returns null.'],
    ['RUN-DOC-STATUS', 'Deliverable completeness gate', 'Review per-case evidence mapping and historical/current fingerprints', 'PARTIAL — DOCX NOT FINALIZED', 'No DOCX status reconciliation performed because the 49-case API run lacks verified per-case evidence; preserve pending/unverified classifications.'],
]
for row in rows:
    ws.append(row)
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions
for col, width in {'A':22,'B':34,'C':58,'D':34,'E':105}.items():
    ws.column_dimensions[col].width = width
for row in ws.iter_rows():
    for cell in row:
        cell.alignment = copy(cell.alignment) + Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[1].height = 32
for i in range(2, ws.max_row+1):
    ws.row_dimensions[i].height = 66
wb.save(path)
print({'path':str(path.resolve()),'sheet':name,'rows':ws.max_row-1,'sheets':len(wb.sheetnames)})
