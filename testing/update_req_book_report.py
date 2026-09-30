"""Merge focused REQ-BOOK evidence into v3 XLSX/DOCX deliverables."""
from pathlib import Path
from collections import Counter
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

root=Path('C:/GuisTee/Project_Y3/Book-a-Badminton/Book-a-Badminton_Court/testing')
src_x=root/'Test_Case_Book-a-Badminton_Court_Results_20260928_v2.xlsx'
src_d=root/'Test_Summary_Report_Book-a-Badminton_Court_v2.docx'
out_x=root/'Test_Case_Book-a-Badminton_Court_Results_20260928_v3_REQ-BOOK.xlsx'
out_d=root/'Test_Summary_Report_Book-a-Badminton-Court_v3_REQ-BOOK.docx'
statuses={
1:'ผ่าน',2:'ผ่าน',3:'ไม่ผ่าน',4:'ไม่ผ่าน',5:'ผ่าน',6:'ผ่าน',7:'ผ่าน',8:'ผ่าน',9:'ตรวจไม่ได้',10:'ผ่าน',11:'ไม่ผ่าน',12:'ผ่าน',13:'ผ่าน',14:'ผ่าน',15:'ผ่าน',16:'ผ่าน',17:'ผ่าน',18:'ผ่าน',19:'ผ่าน',20:'ผ่าน',21:'ผ่าน',22:'ไม่ผ่าน',23:'ผ่าน',24:'ผ่าน',25:'ผ่าน',26:'ผ่าน',27:'ไม่ผ่าน',28:'ไม่ผ่าน',29:'ไม่ผ่าน',30:'ผ่าน',31:'ผ่าน',32:'ผ่าน',33:'ไม่ผ่าน',34:'ตรวจไม่ได้',35:'ตรวจไม่ได้',36:'ไม่ผ่าน',37:'รอทดสอบ',38:'ผ่าน',39:'ผ่าน',40:'ไม่ผ่าน',41:'ผ่าน',42:'ผ่าน'}
status_by_id={f'TC_BOOK_{i:03d}':s for i,s in statuses.items()}
assert Counter(status_by_id.values())=={'ผ่าน':28,'ไม่ผ่าน':10,'ตรวจไม่ได้':3,'รอทดสอบ':1}
actual={
'TC_BOOK_001':'Browser/Chromium + MySQL QA: ยืนยันใน dialog สำเร็จ, toast “จองคอร์ทสำเร็จ”; DB บันทึก 18:00-19:00 สถานะ PENDING.',
'TC_BOOK_002':'Browser/Chromium: dialog แสดง วันที่/เดือน/คอร์ท/เวลา/ชื่อผู้จอง ครบ: 28 | กันยายน | Court 2 | 20:00-21:00 | ชื่อบัญชี QA.',
'TC_BOOK_003':'Browser/Chromium: ปุ่มจริงคือ “ยืนยันจองคอร์ท” ไม่ใช่ข้อความที่ Expected Result กำหนดว่า “จองคอร์ต”.',
'TC_BOOK_005':'API/MySQL QA: วันพรุ่งนี้ตอบ 400 ข้อความ day-by-day policy; จำนวนแถว DB 0→0.',
'TC_BOOK_006':'Browser/Chromium: แสดงวันนี้และคำว่า วันที่; ไม่มี input[type=date] หรือ dropdown ให้เลือกวันอื่น.',
'TC_BOOK_009':'Service/API + MySQL QA: เปลี่ยน fixture เป็น COMPLETED แล้วการจองที่ 2 ตอบ 400 และ DB เหลือ 1 แถว; ยังไม่ได้ทำ lifecycle จอง→check-in→cron COMPLETED ครบจริง.',
'TC_BOOK_011':'API/MySQL QA: startTime=18:30 ตอบ 400 “Invalid booking start time.” ตั้งแต่ validation จึงไม่ได้พิสูจน์การชนแบบช่วงเวลา 18:30-19:30 ตาม Expected Result.',
'TC_BOOK_012':'API/MySQL QA: บันทึกติดกันพอดี 18:00-19:00 และ 19:00-20:00 สำเร็จ 2 รายการ.',
'TC_BOOK_014':'Mocked Date + BookingsService + MySQL QA: Court 3 รอบแรกบันทึก 08:00-09:00 สถานะ PENDING.',
'TC_BOOK_015':'API/DB + Browser: POST Court 4 เวลา 22:00 ได้ 201 และ DB 22:00-23:00; UI แสดง 15 รอบ 08:00-09:00 ถึง 22:00-23:00 ไม่มี 23:00-24:00.',
'TC_BOOK_016':'API/MySQL QA: A จอง 18:00 แล้วยกเลิกทันเวลา; B จองคอร์ต/เวลาเดิมได้ 201; DB เก็บรายการแรก CANCELLED.',
'TC_BOOK_017':'Browser/Chromium: จำลอง availability เก่าให้ B ยังเลือก Court 1 ได้; POST ตอบ 400 และ UI แสดงข้อความ “คอร์ท/เวลาชน”; DB คง 1 รายการ.',
'TC_BOOK_018':'Browser/Chromium: ผู้ใช้มี Court 1 เวลา 18:00 แล้วลอง Court 2 เวลา 20:00; POST 400, UI แสดง “ใช้โควตาประจำวันแล้ว”; DB มี 1 แถว.',
'TC_BOOK_019':'Browser/Chromium: URL วันพรุ่งนี้ถูกพากลับ /booking พร้อม toast “ระบบเปิดให้จองเฉพาะวันนี้เท่านั้น”; ไม่มี POST/แถว DB.',
'TC_BOOK_020':'API/MySQL QA: DB ผูก stu_id ของ token จริง, court=3, วันที่ทดสอบ, 18:00-19:00, PENDING, admin_id=A001.',
'TC_BOOK_021':'Browser/Chromium: รายการใหม่ 18:00/PENDING ปรากฏใน Recent Bookings ภายใน 5 วินาที.',
'TC_BOOK_022':'API/MySQL QA: 10:30, ab:cd, 25:99 ตอบ 400 “Invalid booking start time.” แต่ 24:00 ตอบ “Booking hours are 08:00 to 23:00.” ไม่ตรง Expected Result; ไม่มีแถว DB.',
'TC_BOOK_023':'API/MySQL QA: 20:00:00 และ 20:00 normalize เป็น 20:00-21:00; คอร์ตเดิมชนตอบ 400, คนละคอร์ตบันทึก 201.',
'TC_BOOK_024':'API/MySQL QA: 07:00 และ 23:00 ตอบ 400 พร้อม “Booking hours are 08:00 to 23:00.”; ไม่มีแถว DB.',
'TC_BOOK_025':'API/MySQL QA เวลาไทย 14:06: ส่งรอบ 10:00 ตอบ 400 “This booking round has already ended.”; ไม่มีแถว DB.',
'TC_BOOK_026':'Mocked Date 10:30 + BookingsService + MySQL QA: รอบ 10:00 บันทึก time_in=10:00:00, time_out=11:00:00.',
'TC_BOOK_027':'Browser/Chromium mocked clock: หน้า /booking เดิมกันรอบหมดเวลาได้ แต่เปิด /booking/select-court?time=18:00 โดยตรงหลัง 19:05 ยังแสดง 4 คอร์ตให้เลือก จึงไม่ผ่าน assertion ทั้งเคส.',
'TC_BOOK_028':'Browser/Chromium mocked clock 23:10: ไม่มีรอบที่เลือกได้และปุ่มดำเนินการ disabled แต่ไม่มีข้อความอธิบายหมดเวลา/ปิดรับจองตาม Expected Result.',
'TC_BOOK_030':'Browser + API QA: availability คืน 4 คอร์ตว่าง; UI แสดง 4 cards และเลือกได้ 4 คอร์ตสำหรับ 18:00 ตรงกัน.',
'TC_BOOK_031':'Browser + API/MySQL QA: Court 2 มี PENDING จึง disabled; UI เหลือ 3/4 คอร์ตให้เลือกและตรงกับ availability.',
'TC_BOOK_032':'Browser/Chromium: ไม่มี query หรือขาด time ถูกพากลับ /booking; ไม่ยิง availability ด้วย date ว่าง.',
'TC_BOOK_033':'Browser/Chromium: เมื่อวาน/พรุ่งนี้/2026-02-30 ถูกพากลับ /booking แต่ time=10:30 ของวันนี้ค้างที่ /booking/select-court จึงไม่ผ่าน Expected Result.',
'TC_BOOK_034':'Mocked Date + BookingsService/MySQL: เปิดด้วยวันที่ 28 แล้วเรียกหลังข้ามวัน 29 เวลา 00:01 ถูก backend ปฏิเสธ day-by-day และไม่บันทึก; ยังไม่ได้พิสูจน์ toast/redirect จากหน้า browser ที่เปิดค้าง.',
'TC_BOOK_035':'Browser timezone America/Los_Angeles (วันที่ local 27 vs ไทย 28): หน้า booking ใช้วันที่ไทยและ Dashboard เห็นรายการ 20:00; ยังตรวจรูปแบบวัน/เดือนทุกจุดบน Dashboard ไม่ครบ.',
'TC_BOOK_036':'Browser/Chromium จำลอง availability HTTP 500: UI แสดง toast “Failed to load courts” แต่หน้าแสดง 0 คอร์ทพร้อมใช้และไม่มีปุ่มลองใหม่ จึงไม่ผ่าน.',
'TC_BOOK_038':'Browser/Chromium: ปิด dialog ด้วย “กลับไปเลือก”; POST=0, Booking 0→0, /bookings/me ว่าง และ Court 1 ยังเลือกได้.',
'TC_BOOK_039':'API/MySQL QA: มี COMPLETED เมื่อวาน แล้ววันนี้จอง Court 3 เวลา 18:00 ได้ 201.',
'TC_BOOK_041':'Browser proxy: server commit POST 201 แล้วตัด response ให้ client timeout; UI แจ้ง error, /bookings/me พบ 1 รายการ; retry ตอบ 400 quota; DB มี 1 แถว.',
'TC_BOOK_042':'Mocked transaction + MySQL QA: error ก่อน callback ไม่ทิ้งแถว; คำขอถัดไปบันทึกสำเร็จและไม่พบ lock ค้าง.',
}
evidence='หลักฐาน: qa_book_api_results.json; qa_book_ui_results.json; qa_book_ui_recheck_results.json; qa_book_ui_TC_BOOK_017_TC_BOOK_018_TC_BOOK_019.json; qa_book_ui_TC_BOOK_027_TC_BOOK_028.json; qa_book_ui_TC_BOOK_035.json; qa_book_ui_TC_BOOK_041.json; qa_book_time_results.json'

w=load_workbook(src_x)
s=w['02_REQ-BOOK']; rows={s.cell(r,1).value:r for r in range(3,s.max_row+1)}
assert set(rows)==set(status_by_id)
for cid,status in status_by_id.items():
 r=rows[cid];s.cell(r,7).value=status
 if cid in actual:s.cell(r,6).value=actual[cid]+'\n'+evidence
by_group={sh.title:Counter(sh.cell(r,7).value for r in range(3,sh.max_row+1)) for sh in w.worksheets[1:10]}
counts=Counter({key:sum(g[key] for g in by_group.values()) for key in ['ผ่าน','ไม่ผ่าน','ตรวจไม่ได้','รอทดสอบ']})
assert by_group['02_REQ-BOOK']==Counter({'ผ่าน':28,'ไม่ผ่าน':10,'ตรวจไม่ได้':3,'รอทดสอบ':1})
assert counts==Counter({'ผ่าน':72,'ไม่ผ่าน':22,'ตรวจไม่ได้':38,'รอทดสอบ':139})
summary=w['00_Summary']
summary.cell(37,1).value='ผลทดสอบจริง 28/09/2569: เพิ่มรอบโฟกัส REQ-BOOK บน API/MySQL, Chromium และ BookingsService+mocked Date; QA volume แยก; เฉพาะเคสที่ครบทุก assertion จึงนับผ่าน/ไม่ผ่าน'
for col,key in [(2,'ผ่าน'),(4,'ไม่ผ่าน'),(6,'ตรวจไม่ได้'),(8,'รอทดสอบ')]:summary.cell(38,col).value=counts[key]
summary.cell(39,1).value='ข้อจำกัด: REQ-BOOK ยังเหลือ TC_BOOK_037 รอทดสอบ และ TC_BOOK_009/034/035 ตรวจบางส่วน; กล้อง, โหลดสูง, production HTTPS, หลาย instance และนโยบายอื่นยังอยู่นอกการรันนี้'
summary.cell(40,1).value='หลักฐาน REQ-BOOK: testing/evidence/qa_book_*.json; runners: testing/run_qa_book_api.cjs, run_qa_book_ui.cjs, run_qa_book_time.cjs'
summary.append(['รอบ REQ-BOOK 28/09/2569: 42 เคส — ผ่าน 28, ไม่ผ่าน 10, ตรวจบางส่วน 3, รอทดสอบ 1; ผลรวมทั้ง workbook ผ่าน 72, ไม่ผ่าน 22, ตรวจบางส่วน 38, รอทดสอบ 139 จาก 271'])
ex=w.create_sheet('14_REQ-BOOK_Execution')
ex.append(['Test Case ID','Status','Actual Result','Evidence Set','Build'])
for i in range(1,43):
 cid=f'TC_BOOK_{i:03d}';ex.append([cid,status_by_id[cid],s.cell(rows[cid],6).value,'testing/evidence/qa_book_*.json','feature-Docx 37aea29 / code 581a654 / cases ed3bdca'])
for c in ex[1]:c.fill=PatternFill('solid',fgColor='DCE6F1');c.font=Font(name='Tahoma',size=10,bold=True)
for row in ex.iter_rows(min_row=2):
 for c in row:c.font=Font(name='Tahoma',size=9);c.alignment=Alignment(wrap_text=True,vertical='top')
 ex.row_dimensions[row[0].row].height=76
for col,width in {'A':20,'B':16,'C':100,'D':42,'E':48}.items():ex.column_dimensions[col].width=width
ex.freeze_panes='A2';ex.auto_filter.ref=f'A1:E{ex.max_row}'
w.calculation.fullCalcOnLoad=True
w.save(out_x)

d=Document(src_d)
def setp(i,text):p=d.paragraphs[i];p.clear();p.add_run(text)
setp(7,f"สถานะความพร้อมภาพรวม: ยังไม่พร้อม Go-Live — ผ่าน {counts['ผ่าน']}, ไม่ผ่าน {counts['ไม่ผ่าน']}, ตรวจบางส่วน/ตรวจไม่ได้ {counts['ตรวจไม่ได้']}, รอทดสอบ {counts['รอทดสอบ']} จาก 271 เคส; รอบ REQ-BOOK พบความต่าง contract และช่องว่าง validation/error state เพิ่ม")
resolved=counts['ผ่าน']+counts['ไม่ผ่าน']
setp(9,f"Pass Rate ของเคสที่ตัดสินผลครบ = ผ่าน {counts['ผ่าน']} ÷ ({counts['ผ่าน']}+{counts['ไม่ผ่าน']}) = {counts['ผ่าน']/resolved*100:.1f}% (เฉพาะ {resolved} เคสที่ตัดสินผลแล้ว ไม่ใช่อัตราผ่านของ 271 เคส)")
setp(10,'ความเสี่ยงหลัก: ช่องโหว่บัญชี/ข้อมูล hash เดิม; direct URL รอบหมดเวลายังเลือกคอร์ตได้; query เวลา 10:30 ไม่ถูกปฏิเสธที่หน้า select-court; availability 500 ไม่มี retry และดูเหมือน 0 คอร์ตว่าง; ข้อความปุ่ม/validation บางกรณีไม่ตรงข้อกำหนด')
setp(13,'รอบเดิมทดสอบ Jest, API/MySQL, service/cron และ Chromium; รอบนี้โฟกัส REQ-BOOK ทั้ง API/DB, flow browser จริง, mocked timezone/clock และ BookingsService + mocked Date กับ MySQL QA แยก พร้อมตรวจ cleanup fixture')
setp(15,f"เคสที่ยังไม่จบ {counts['รอทดสอบ']+counts['ตรวจไม่ได้']} รายการ; ใน REQ-BOOK เหลือ TC_BOOK_037 รอทดสอบ และ TC_BOOK_009/034/035 ตรวจได้บางส่วน ส่วนกล้องจริง หลาย instance โหลดสูง Safari/iOS production HTTPS และ backup/restore ยังต้องทดสอบเพิ่ม")
setp(20,'คำสั่งรอบ REQ-BOOK: node testing/run_qa_book_api.cjs; node testing/run_qa_book_ui.cjs; node testing/run_qa_book_time.cjs โดยมี QA_ISOLATED_DB guard และตรวจ Docker volume badminton_qa_20260928_mysql_data; Chromium headless; mocked Date เฉพาะเคสเวลา; fixture ถูกลบหลังรัน')
setp(27,'Defect Summary ด้านล่างนับเป็น 14 กลุ่มปัญหา/contract mismatch ไม่ใช่จำนวน Test Case ที่ล้ม; รอบ REQ-BOOK เพิ่ม 5 กลุ่มและยังไม่มีหลักฐานว่าแก้ไขแล้ว')
setp(29,'รอบ REQ-BOOK 28/09/2569: 42 เคส — ผ่าน 28, ไม่ผ่าน 10, ตรวจบางส่วน 3, รอทดสอบ 1. ปัญหาใหม่หลัก: ข้อความปุ่ม U04 ไม่ตรง, 18:30/24:00 ไม่ตรง Expected, invalid/expired query, ไม่มีคำอธิบายหลัง 23:00 และ availability 500 ไม่มี retry; ดู testing/evidence/qa_book_*.json')
setp(32,'ความเสี่ยง: ข้อบกพร่องเดิมด้านบัญชี/ข้อมูลลับยังอยู่; REQ-BOOK direct URL อาจนำไปสู่ขั้นเลือกคอร์ตด้วยรอบผิดกติกา, error availability ทำให้เข้าใจว่าไม่มีคอร์ตว่าง และ contract UI/validation ไม่ตรงตารางทดสอบ')
setp(33,'แนวทางลดความเสี่ยง: แก้ validation date/time ที่ select-court และตรวจซ้ำก่อนแสดงคอร์ต; เพิ่ม error state + retry; ตกลงข้อความปุ่มและ semantics 18:30/24:00 แล้วแก้โค้ดหรือ Expected Result ให้ตรง; เพิ่ม regression tests ก่อน Go-Live')
setp(37,'Detailed Test Cases (รอบ REQ-BOOK): testing/Test_Case_Book-a-Badminton_Court_Results_20260928_v3_REQ-BOOK.xlsx — source cases ed3bdca; runner/report branch 37aea29; code behavior 581a654')
setp(38,'หลักฐานรอบ REQ-BOOK: testing/evidence/qa_book_api_results.json, qa_book_ui_results.json, qa_book_ui_recheck_results.json, qa_book_ui_TC_BOOK_*.json และ qa_book_time_results.json; runners testing/run_qa_book_*.cjs; ใช้ MySQL QA แยกและ Chromium headless')
# Overall table
for i,(name,stats) in enumerate(by_group.items(),1):
 rr=stats['ผ่าน']+stats['ไม่ผ่าน'];vals=[name,sum(stats.values()),stats['ผ่าน'],stats['ไม่ผ่าน'],stats['ตรวจไม่ได้'],stats['รอทดสอบ'],f"{stats['ผ่าน']/rr*100:.1f}%" if rr else 'N/A']
 for c,v in enumerate(vals):d.tables[0].cell(i,c).text=str(v)
for c,v in enumerate(['รวมทั้งสิ้น (Total)',271,counts['ผ่าน'],counts['ไม่ผ่าน'],counts['ตรวจไม่ได้'],counts['รอทดสอบ'],f"{counts['ผ่าน']/resolved*100:.1f}%"]):d.tables[0].cell(10,c).text=str(v)
d.tables[1].cell(3,1).text=f"{counts['ผ่าน']} ผ่าน / {counts['ไม่ผ่าน']} ไม่ผ่าน / {counts['ตรวจไม่ได้']} ตรวจไม่ได้ / {counts['รอทดสอบ']} รอทดสอบ"
d.tables[1].cell(3,2).text='รวมหลักฐานเดิมและรอบโฟกัส REQ-BOOK; Jest ไม่ถูกนับเพิ่ม'
row=d.tables[1].add_row().cells
for c,v in enumerate(['REQ-BOOK focused','42 เคส: 28 ผ่าน, 10 ไม่ผ่าน, 3 ตรวจบางส่วน, 1 รอ','API/MySQL + Chromium + mocked Date; qa_book_*.json; fixture cleanup verified']):row[c].text=v
# Severity: prior 9 + 5 new = 14
sev=[('Critical / Blocker',1,0,1,0),('High',2,0,2,0),('Medium',8,0,8,0),('Low',3,0,3,0),('รวมทั้งสิ้น',14,0,14,0)]
for r,vals in enumerate(sev,1):
 for c,v in enumerate(vals):d.tables[2].cell(r,c).text=str(v)
new_defects=[
('QA-BOOK-01','Low','TC_BOOK_003: ปุ่มจริง “ยืนยันจองคอร์ท” ไม่ตรง Expected “จองคอร์ต”','Open / contract mismatch','ยืนยันข้อความ U04 แล้วแก้ UI หรือ Expected Result ให้ตรง'),
('QA-BOOK-02','Medium','TC_BOOK_011/022: 18:30 ถูก validation ปฏิเสธก่อน overlap; 24:00 ได้ข้อความ hours ไม่ใช่ invalid start time','Open / contract mismatch','กำหนด oracle สำหรับครึ่งชั่วโมง/24:00 และจัดลำดับ validation ให้ชัด'),
('QA-BOOK-03','Medium','TC_BOOK_027/033: direct URL รอบหมดเวลาหรือ time=10:30 ยังแสดงหน้าเลือกคอร์ต','Open','validate isBookingSlotSelectable ก่อนโหลด/แสดงคอร์ตและก่อนเปิด dialog'),
('QA-BOOK-04','Low','TC_BOOK_028: หลัง 23:00 ไม่มีรอบให้เลือกแต่ไม่มีคำอธิบายชัดเจน','Open','เพิ่ม empty state ปิดรับจอง/หมดเวลา พร้อมทางกลับ'),
('QA-BOOK-05','Medium','TC_BOOK_036: availability HTTP 500 แสดง 0 คอร์ตพร้อมใช้และไม่มีปุ่ม retry','Open','แยก error state จาก empty state และเพิ่ม retry ที่เรียก API ใหม่'),
]
t=d.tables[3]
for entry in new_defects:
 row=t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 for c,v in enumerate(entry):row.cells[c].text=v
 for cell in row.cells:
  for p in cell.paragraphs:
   p.paragraph_format.space_after=Pt(0)
   for run in p.runs:run.font.size=Pt(8)
# Explicit TH SarabunPSK in styles and runs.
def font(f):
 f.name='TH SarabunPSK';rpr=f._element.get_or_add_rPr() if hasattr(f._element,'get_or_add_rPr') else f._element;rfonts=rpr.rFonts
 if rfonts is None:rfonts=OxmlElement('w:rFonts');rpr.insert(0,rfonts)
 for a in ('ascii','hAnsi','eastAsia','cs'):rfonts.set(qn('w:'+a),'TH SarabunPSK')
for style in d.styles:
 if style.type in (1,2):font(style.font)
for p in d.paragraphs:
 for run in p.runs:font(run.font)
for table in d.tables:
 for row in table.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for run in p.runs:font(run.font)
d.save(out_d)
# Reopen and verify exact totals/status/font.
check=load_workbook(out_x,read_only=True);cs=check['02_REQ-BOOK']
assert Counter(cs.cell(r,7).value for r in range(3,cs.max_row+1))==by_group['02_REQ-BOOK']
assert sum(Counter(check[sh.title].cell(r,7).value for sh in check.worksheets[1:10] for r in range(3,sh.max_row+1)).values())==271
rd=Document(out_d);assert rd.tables[0].cell(10,2).text=='72';assert rd.tables[2].cell(5,1).text=='14';assert rd.styles['Normal'].font.name=='TH SarabunPSK'
print(json.dumps({'overall':counts,'req_book':by_group['02_REQ-BOOK'],'xlsx':str(out_x),'docx':str(out_d),'defects':14},ensure_ascii=False))
