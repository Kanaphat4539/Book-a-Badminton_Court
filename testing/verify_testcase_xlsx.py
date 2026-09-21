# -*- coding: utf-8 -*-
"""Verify the generated test-case workbook against test.md, the backup and the mapping sheet."""
import os
import re
from openpyxl import load_workbook

ROOT = r"C:/GuisTee/Project_Y3/Book-a-Badminton/Book-a-Badminton_Court"
NEW = ROOT + "/testing/Test_Case_Book-a-Badminton_Court.xlsx"
OLD = ROOT + "/testing/Test_Case_Book-a-Badminton_Court_backup_e9acf4a.xlsx"
TESTMD = os.path.join(os.path.dirname(os.path.abspath(__file__)), "verify_source_test.md")
if not os.path.exists(TESTMD):  # fall back to the working copy used while building
    TESTMD = r"C:/Users/TEE/AppData/Local/hermes/cache/scratch/dev_test.md"

HEAD = ["Test Case ID", "Test Scenario", "Test Steps", "Test Data", "Expected Result",
        "Actual Result", "Status"]
SKIP = {"00_Summary", "10_Rules_dev", "11_Mapping"}
# prefixes of test.md categories 3-20 (one sheet per group); the rest are setup/smoke/perf/deploy/GAP
REQ_PREFIX = {"REG", "AUT", "AVA", "BKG", "CAN", "QR", "TIM", "BAN", "CRN", "ADM", "USR", "RST",
              "NTF", "NWS", "DSH", "UI", "ACL", "SEC", "RACE", "DATA", "API"}

wb = load_workbook(NEW)
old = load_workbook(OLD)

# --- 1. the removed team table is not referenced anywhere -------------------
hits = [(n, c.coordinate) for n in wb.sheetnames for row in wb[n].iter_rows() for c in row
        if isinstance(c.value, str) and "Test-table" in c.value]
print("1) references to the removed team table:", hits or "none")
assert not hits

# --- 2. case sheets ---------------------------------------------------------
ids, counts, deltas = [], [], {}
for name in wb.sheetnames:
    if name in SKIP:
        continue
    ws = wb[name]
    assert [ws.cell(2, i).value for i in range(1, 8)] == HEAD, name
    n = ws.max_row - 2
    counts.append((name, n))
    ov = old[name].max_row - 2
    deltas[name] = n - ov
    assert n >= ov, (name, n, ov)
    dv = [str(d.sqref) for d in ws.data_validations.dataValidation]
    assert any(s.startswith(f"G3:G{2+n}") for s in dv), (name, dv)
    for r in range(3, ws.max_row + 1):
        ids.append(ws.cell(r, 1).value)
        for c in (2, 3, 4, 5):
            assert str(ws.cell(r, c).value or "").strip(), (name, r, c)
        assert ws.cell(r, 7).value == "รอทดสอบ", (name, r)
    print(f"   {name}: {n} cases (+{deltas[name]}), dropdown {dv[0]}, freeze {ws.freeze_panes}")
total = sum(c for _, c in counts)
print(f"2) total cases {total} | unique ids {len(set(ids))} | new vs backup {sum(deltas.values())}")
assert len(set(ids)) == len(ids), "duplicate Test Case ID"
assert total == len(ids)
assert total == sum(old[n].max_row - 2 for n, _ in counts) + sum(deltas.values())

# --- 3. summary count column ------------------------------------------------
s = wb["00_Summary"]
sc = dict(counts)
for r in range(12, 22):
    nm, n = s.cell(r, 4).value, s.cell(r, 5).value
    if nm in sc:
        assert n == sc[nm], (nm, n, sc[nm])
    elif nm == "รวม":
        assert n == total, n
print("3) summary count column matches all sheets:", True)

# --- 4. mapping sheet -------------------------------------------------------
m = wb["11_Mapping"]
assert m.max_column == 5, m.max_column
assert m["A1"].value.startswith("ความครบของ Test Case"), m["A1"].value
case_ids = set(ids)
covered_all, missing_all, nosheet_total = [], [], 0
for r in range(3, m.max_row + 1):
    a, b, tids, cov, miss = (m.cell(r, i).value for i in range(1, 6))
    declared = int(re.search(r"\((\d+) รายการ\)", tids).group(1))
    if str(miss).startswith("ไม่สร้างชีต Test Case:"):
        assert cov == "—", (r, "no-sheet row must not claim coverage")
        nosheet_total += declared
        continue
    pairs = re.findall(r"([A-Z]{2,4}-\d{2}) → ([^;]+)", cov or "")
    ids_here = [p[0] for p in pairs]
    for tid, tcs in pairs:
        for tc in re.findall(r"TC_[A-Z]+_\d{3}", tcs):
            assert tc in case_ids, (r, tid, tc, "TC referenced by mapping but missing from sheets")
    assert len(set(ids_here)) == len(ids_here), (r, "duplicate id inside a row")
    assert len(ids_here) == declared, (r, declared, len(ids_here))
    assert str(miss).startswith("ครบทั้งหมด:"), (r, miss[:40])
    covered_all += ids_here
    missing_all += re.findall(r"[A-Z]{2,4}-\d{2}", miss.split("—")[0])
print(f"4) mapping: covered {len(covered_all)}, missing {len(missing_all)}, no-sheet {nosheet_total}")
assert not missing_all
assert len(set(covered_all)) == len(covered_all), "an ID is claimed by more than one group"

# --- 5. independent re-derivation from test.md ------------------------------
md_ids = re.findall(r"^- \[[ x]\] ([A-Z]{2,4}-\d{2})", open(TESTMD, encoding="utf-8").read(), re.M)
md_req = {i for i in md_ids if i.split("-")[0] in REQ_PREFIX}
print(f"5) test.md: {len(md_ids)} ids total, {len(md_req)} in categories 3-20")
assert len(md_ids) == len(set(md_ids))
assert md_req == set(covered_all), ("coverage mismatch", sorted(md_req - set(covered_all))[:10],
                                    sorted(set(covered_all) - md_req)[:10])
print("   every test.md id in categories 3-20 is claimed covered: True")

# --- 6. notes ---------------------------------------------------------------
notes = [s.cell(r, 1).value for r in range(1, s.max_row + 1) if isinstance(s.cell(r, 1).value, str)]
note12 = [n for n in notes if n.startswith("12.")][0]
assert f"ทั้งหมด {len(md_ids)} รายการ" in note12, note12
assert f"หมวด 3-20 รวม {len(md_req)} รายการ" in note12, note12
assert f"ครอบแล้ว {len(covered_all)} รายการ" in note12, note12
assert "ไม่มี ID ที่ยังขาดเคสเฉพาะ (0 รายการ)" in note12, note12
assert f"อีก {nosheet_total} รายการ" in note12, note12
assert len(md_req) + nosheet_total == len(md_ids)
assert any(n.startswith("13.") and "105 รายการ" in n for n in notes), "note 13 missing"
print("6) summary note 12 matches test.md, note 13 present: True")

# --- 7. existing cases were not edited, only appended -----------------------
edited = []
for n in wb.sheetnames:
    for r in range(1, old[n].max_row + 1):
        for c in range(1, max(old[n].max_column, wb[n].max_column) + 1):
            a, b = old[n].cell(r, c).value, wb[n].cell(r, c).value
            if a != b:
                edited.append((n, old[n].cell(r, c).coordinate, a, b))
# allowed: reference sheets (version/notes/coverage), and the commit marker e9acf4a -> 0fcde23
bad = [e for e in edited
       if e[0] not in {"00_Summary", "10_Rules_dev", "11_Mapping"}
       and not (isinstance(e[2], str) and isinstance(e[3], str)
                and "dev e9acf4a" in e[2] and e[3] == e[2].replace("dev e9acf4a", "dev 0fcde23"))]
print(f"7) existing rows: {len(edited)} changed cells in {sorted({e[0] for e in edited})}"
      f" | unexpected: {len(bad)} {bad[:3]}")
assert not edited or all(e[0] in {"00_Summary", "10_Rules_dev", "11_Mapping"}
                         or "dev e9acf4a" in str(e[2]) for e in edited)
assert not bad, bad[:5]
print("VERIFY OK")
